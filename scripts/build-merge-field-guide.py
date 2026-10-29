"""Build the customer Word field guide from the live PHP field catalog.

Run from the repository root with python-docx installed and PHP on PATH.
"""
import json
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
php = r'''require 'vendor/autoload.php'; echo json_encode(array_map(fn($m) => App\Http\Controllers\Api\V1\PrintTemplateController::fields($m), array_combine(array_keys(App\Http\Controllers\Api\V1\PrintTemplateController::modules()), array_keys(App\Http\Controllers\Api\V1\PrintTemplateController::modules()))), JSON_UNESCAPED_UNICODE);'''
catalog = json.loads(subprocess.check_output(['php', '-r', php], cwd=ROOT))
modules = dict(quotation='Báo giá', sales_order='Đơn bán hàng', goods_issue='Phiếu xuất kho', contract='Hợp đồng', delivery_acceptance='Biên bản nghiệm thu')
sets = {m: dict((f['key'], f['label']) for f in fs) for m, fs in catalog.items()}
common = set.intersection(*(set(fs) for fs in sets.values()))
examples = {
    'ty_le_dat_coc':'Theo tỷ lệ đã chọn, ví dụ 30%', 'ty_le_con_lai':'Tỷ lệ còn lại, ví dụ 70%',
    'tien_dat_coc':'Tiền cọc dự kiến theo lựa chọn', 'tien_con_lai':'Số tiền còn lại dự kiến',
    'nhan_thanh_toan_dau':'Nhãn tự đổi theo cách thanh toán', 'gia_tri_thanh_toan_dau':'Tiền cọc, thanh toán 100% hoặc nội dung riêng',
    'nhan_thanh_toan_con_lai':'Nhãn số tiền còn lại; trống nếu không đặt cọc', 'gia_tri_thanh_toan_con_lai':'Tiền còn lại; trống nếu không đặt cọc',
    'han_dat_coc':'Ngày đến hạn cọc', 'han_thanh_toan_con_lai':'Ngày đến hạn khoản còn lại', 'so_ngay_cong_no':'Số ngày công nợ đã chọn',
    'tien_dat_coc_70': '70% tổng báo giá, làm tròn đến đồng; không phải tiền đã thu',
    'tien_con_lai_30': 'Tổng báo giá trừ cọc dự kiến 70%; không phải số dư công nợ',
    'ma_chung_tu': 'Theo mã chứng từ đang xuất', 'ngay_chung_tu': '27/10/2026',
    'ten_khach_hang': 'Công ty Minh An', 'ma_khach_hang': 'CUS-00001',
    'nguoi_lien_he': 'Nguyễn An', 'so_dien_thoai': '0901234567',
    'dia_chi_giao_hang': '12 Nguyễn Huệ, TP.HCM', 'stt': '1',
    'ma_hang': 'SKU-00001', 'ten_hang': 'Sản phẩm A', 'dvt': 'Cái',
    'so_luong': '2', 'bang_dong_hang': 'Danh sách hàng dạng văn bản',
    'nguoi_lap': 'Tên người lập chứng từ', 'trang_thai': 'Theo trạng thái chứng từ',
    'email': 'lienhe@example.com', 'ten_cong_ty': 'Công ty khách hàng',
    'ma_so_thue': 'Mã số thuế khách hàng', 'dia_chi_cong_ty': 'Địa chỉ hóa đơn khách hàng',
    'nguoi_dai_dien_phap_luat': 'Nguyễn An', 'chuc_vu_dai_dien': 'Giám đốc',
    'don_gia': '100.000 đ', 'vat_percent': '10%', 'vat_dong': '20.000 đ',
    'thanh_tien': 'Giá trị của dòng hàng', 'tong_cong': '220.000 đ',
    'tong_tien_bang_chu': 'Hai Trăm Hai Mươi Nghìn đồng',
    'tam_tinh': '200.000 đ', 'chiet_khau': '0 đ', 'vat': '20.000 đ',
    'ma_bao_gia': 'Mã báo giá liên quan', 'ngay_bao_gia': '27/10/2026',
    'ngay_het_hieu_luc': '30/11/2026', 'dieu_khoan_thanh_toan': 'Theo thỏa thuận trên báo giá',
    'dieu_khoan_giao_hang': 'Theo điều khoản giao hàng', 'ma_don_ban': 'Mã đơn bán liên quan',
    'hinh_thuc_giao_hang': 'Giao tận nơi', 'ghi_chu_giao_hang': 'Giao trong giờ hành chính',
    'kho_xuat': 'Kho trung tâm', 'ly_do_xuat': 'Xuất bán hàng',
    'dia_diem_giao_hang': 'Địa điểm giao trên phiếu', 'chung_tu_goc': 'Chứng từ tham chiếu',
    'co_tru_ton': 'Theo thiết lập trừ tồn của phiếu', 'ma_hop_dong': 'HĐ-001',
    'ten_hop_dong': 'Hợp đồng cung cấp hàng hóa', 'ngay_hop_dong': '27/10/2026',
    'ngay_hieu_luc': '28/10/2026', 'ngay_het_han': '31/12/2026',
    'gia_tri_hop_dong': 'Giá trị đã lưu trên hợp đồng', 'gia_tri_bang_chu': 'Giá trị hợp đồng bằng chữ',
    'ma_don_hang': 'Đơn hàng tham chiếu', 'dieu_khoan': 'Điều khoản đã lưu',
    'moc_thanh_toan': 'Các mốc đã lập trong hợp đồng', 'ngay_giao_hang': '28/10/2026',
    'ngay_nghiem_thu': '29/10/2026', 'nguoi_nghiem_thu': 'Trần Bình',
    'chuc_vu_nghiem_thu': 'Trưởng bộ phận', 'ket_qua_nghiem_thu': 'Kết quả đã ghi nhận',
    'kiem_tra_so_luong': 'Kết quả kiểm tra số lượng', 'kiem_tra_tinh_trang': 'Kết quả kiểm tra tình trạng',
    'kiem_tra_ho_so': 'Kết quả kiểm tra hồ sơ', 'noi_dung_nghiem_thu': 'Nội dung đã ghi nhận',
    'sai_lech_nghiem_thu': 'Sai lệch và yêu cầu đã ghi nhận',
}
assert all(k in examples for fs in sets.values() for k in fs), 'Add an explanation for each new field'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.5)
sec.left_margin = sec.right_margin = Cm(1.4)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(17)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('VK CRM  •  Bộ trường trộn  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def table(rows, widths=(4.8, 7.6, 5.8), headers=('Thông tin', 'Trường để sao chép', 'Diễn giải hoặc ví dụ')):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c, text, width in zip(t.rows[0].cells, headers, widths):
        c.text = text; c.width = Cm(width)
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EDF3FB'); c._tc.get_or_add_tcPr().append(shade)
        for r in c.paragraphs[0].runs: r.bold = True
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = t.add_row().cells
        for i, (c, value, width) in enumerate(zip(cells, values, widths)):
            c.width = Cm(width); c.text = value
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                for r in para.runs:
                    r.font.size = Pt(9)
                    if i == 1: r.font.name = 'Consolas'
        no_split = OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
    t._tbl.tblPr.append(borders)
    return t

doc.add_heading('Bộ trường trộn', 0)
p('VK CRM  |  Tài liệu dùng khi thiết kế mẫu Word')
p('Một file tra cứu cho Báo giá, Đơn bán hàng, Phiếu xuất kho, Hợp đồng và Biên bản nghiệm thu. Bạn có thể tự sắp xếp chữ, bảng, logo và vùng ký tên trong mẫu của mình.')
doc.add_heading('Cách sử dụng', 1)
for text in [
    'Mở một file Word riêng hoặc tải mẫu tham khảo trong trang Mẫu in.',
    'Chọn loại chứng từ. Sao chép trường trong phần dùng chung và phần riêng của loại đó vào vị trí cần hiển thị.',
    'Giữ nguyên hai cặp ngoặc nhọn và mã bên trong, ví dụ {{ten_khach_hang}}. Không dịch mã trường sang tiếng Việt.',
    'Lưu thành .docx, vào Mẫu in → chọn đúng loại → Tải mẫu Word lên. Sau đó xuất thử từ chứng từ có dữ liệu để kiểm tra.',
]: p(text, 'List Number')
p('Đây là tài liệu tra cứu. Hãy sao chép các trường cần dùng vào mẫu riêng; không tải nguyên bộ tài liệu này lên làm mẫu in vì có các trường thuộc nhiều loại chứng từ.')
doc.add_heading('Tra theo nội dung', 1)
for text in ['Trường dùng chung cho cả 5 loại chứng từ', 'Báo giá', 'Đơn bán hàng', 'Phiếu xuất kho', 'Hợp đồng', 'Biên bản nghiệm thu', 'Cách tạo bảng hàng và kiểm tra mẫu']:
    p(text, 'List Bullet')
p('Dữ liệu minh họa trong tài liệu chỉ để giải thích. Khi xuất Word, hệ thống lấy dữ liệu của chứng từ đang chọn. Thông tin chưa được nhập có thể để trống.')
doc.add_page_break()
doc.add_heading('Trường dùng chung', 1)
p('Dùng được cho cả 5 loại chứng từ. Các phần tiếp theo chỉ liệt kê trường bổ sung cho từng loại.')
table([(label, '{{'+key+'}}', examples[key]) for key,label in sets['quotation'].items() if key in common])
p('Tên công ty, mã số thuế và người đại diện trong các phần Báo giá và Hợp đồng là thông tin khách hàng. Thông tin công ty bán hàng, logo và chữ ký cố định có thể tự đặt trực tiếp trong Word.')
for module, title in modules.items():
    doc.add_page_break()
    doc.add_heading(title, 1)
    p('Phạm vi sử dụng: mẫu '+title.lower()+'. Có thể kết hợp các trường dưới đây với toàn bộ trường dùng chung ở phần trước.')
    table([(label, '{{'+key+'}}', examples[key]) for key,label in sets[module].items() if key not in common])
    if module == 'contract': p('Dòng hàng của hợp đồng lấy từ đơn hàng liên kết; nếu không có thì lấy từ báo giá liên kết. Hợp đồng chưa liên kết có thể không có dòng hàng.')
    if module == 'delivery_acceptance': p('Dòng hàng lấy từ phiếu xuất kho liên kết, hoặc đơn bán hàng khi không có phiếu xuất kho. Kết quả kiểm tra lấy từ thông tin nghiệm thu đã lưu.')
doc.add_page_break()
doc.add_heading('Tạo bảng hàng trong Word', 1)
p('Chèn một bảng Word gồm dòng tiêu đề và đúng một dòng mẫu dữ liệu. Đặt các trường dòng hàng trong dòng mẫu; hệ thống tự lặp dòng đó theo số mặt hàng của chứng từ. Giữ dòng tổng tiền bên ngoài dòng mẫu.')
table([('{{stt}}', '{{ten_hang}}', '{{so_luong}}')], widths=(2, 11.2, 5), headers=('STT', 'Tên hàng', 'Số lượng'))
p('Trường dòng hàng dùng chung: {{stt}}, {{ma_hang}}, {{ten_hang}}, {{dvt}}, {{so_luong}}. Báo giá và Đơn bán hàng còn có {{don_gia}}, {{vat_percent}}, {{vat_dong}}, {{thanh_tien}}. Không đặt trường dòng hàng vào dòng tiêu đề hoặc dòng tổng cộng.')
p('Nếu chỉ muốn danh sách hàng dạng văn bản, dùng {{bang_dong_hang}} trong một đoạn riêng. Muốn trình bày bảng đẹp và rõ cột, dùng bảng Word như hướng dẫn trên.')
doc.add_heading('Kiểm tra trước khi sử dụng', 1)
for text in ['Chọn đúng loại mẫu trước khi tải lên. Trường của Báo giá không tự dùng được trong Phiếu xuất kho.', 'Xuất thử chứng từ có từ hai mặt hàng để kiểm tra dòng tự lặp và dòng tổng.', 'Kiểm tra thông tin khách hàng, địa chỉ, ngày và tiền. Nếu trống, kiểm tra dữ liệu gốc và liên kết chứng từ.', 'Logo, thông tin bên bán và bố cục ký tên là nội dung bạn tự thiết kế trong file Word. Chỉ các mã được liệt kê trong tài liệu này mới được hỗ trợ.']:
    p(text, 'List Bullet')
doc.add_page_break()
doc.add_heading('Đối chiếu bộ biểu mẫu Sơn Tiên', 1)
p('Phần này bao quát các nhóm thông tin cần điền trong 9 tài liệu của bộ THỦ TỤC MẪU SƠN TIÊN. Các mã có hai cặp ngoặc nhọn là trường đang hỗ trợ ở đúng loại chứng từ đã ghi trong các phần trước. Các thông tin ghi Cần bổ sung chưa có mã trường trộn hoạt động; chưa thể tự điền chỉ bằng việc đưa tên thông tin vào Word.')
p('Nội dung điều khoản, quy định hủy đổi, logo, tiêu đề và khung ký có thể giữ nguyên trong mẫu Word sau khi doanh nghiệp xác nhận. Các tỷ lệ cọc, chính sách miễn và thời hạn trong tài liệu mẫu không phải quy tắc mặc định cho mọi giao dịch.')
sections = [
 ('Báo giá dịch vụ', [
  ('Số và ngày báo giá', '{{ma_bao_gia}} / {{ngay_bao_gia}}', 'Đã hỗ trợ trong Báo giá'),
  ('Đơn vị đặt dịch vụ và liên hệ', '{{ten_khach_hang}} / {{nguoi_lien_he}} / {{so_dien_thoai}} / {{email}}', 'Đã hỗ trợ trong Báo giá'),
  ('Địa chỉ bên đặt và đại diện', '{{dia_chi_cong_ty}} / {{nguoi_dai_dien_phap_luat}} / {{chuc_vu_dai_dien}}', 'Đã hỗ trợ trong Báo giá'),
  ('Hạng mục và chi phí', '{{stt}} / {{ten_hang}} / {{dvt}} / {{so_luong}} / {{don_gia}} / {{thanh_tien}}', 'Đã hỗ trợ trong Báo giá'),
  ('Tổng tiền và bằng chữ', '{{tong_cong}} / {{tong_tien_bang_chu}}', 'Đã hỗ trợ trong Báo giá'),
  ('Đặt cọc và còn lại', '{{nhan_thanh_toan_dau}} / {{gia_tri_thanh_toan_dau}} / {{nhan_thanh_toan_con_lai}} / {{gia_tri_thanh_toan_con_lai}}', 'Đã hỗ trợ theo cách thanh toán lưu trong báo giá; không phải tiền thực thu'),
  ('Bên cung cấp và chuyển khoản', 'Tên, địa chỉ, điện thoại, email, website, đại diện; ngân hàng, tài khoản, chủ tài khoản, nội dung chuyển khoản', 'Giữ cố định trong Word hoặc bổ sung cấu hình bên bán; không dùng thông tin công ty khách hàng thay thế'),
 ]),
 ('Phiếu xác nhận dịch vụ đoàn', [
  ('Thông tin phiếu', 'Số phiếu, ngày lập, báo giá và hợp đồng liên quan, người xác nhận', 'Cần bổ sung loại chứng từ và trường riêng'),
  ('Bên đặt và bên cung cấp', 'Đơn vị đặt, trường, đại diện, điện thoại, email, địa chỉ; thông tin bên cung cấp', 'Tái sử dụng hồ sơ khách và cấu hình công ty khi nối loại phiếu mới'),
  ('Lịch và phân loại đoàn', 'Ngày tham quan, dịch vụ, đối tượng, chính sách áp dụng', 'Cần nối lịch đoàn và lựa chọn chính sách'),
  ('Vé theo đối tượng', 'Người lớn, trẻ em, hướng dẫn viên; số đăng ký, số miễn, số thanh toán, đơn giá, thành tiền', 'Cần bổ sung bảng vé và tính số lượng thanh toán; FOC là suất miễn phí'),
  ('Ăn uống', 'Thực đơn đính kèm, tên thực đơn, số lượng, đơn vị tính, đơn giá, thành tiền', 'Cần bổ sung bảng ăn uống và tệp đính kèm'),
  ('Dịch vụ khác', 'Tên dịch vụ, số lượng, đơn vị tính, đơn giá, thành tiền', 'Cần bổ sung nhóm dịch vụ khác'),
  ('Tổng và quyết toán', 'Tổng vé, tổng ăn uống, tổng dịch vụ khác, tổng chung, bằng chữ, cọc, còn lại, phát sinh', 'Cần bổ sung tổng theo nhóm và liên kết thu tiền'),
  ('Điều chỉnh và xác nhận', 'Ngày thay đổi, hủy, đổi lịch, tăng giảm số lượng; xác nhận bên đặt và kinh doanh', 'Cần bổ sung lịch sử phiên bản và xác nhận; nội dung quy định có thể giữ cố định'),
 ]),
 ('BEO triển khai phục vụ đoàn', [
  ('Đầu mối', 'Tên công ty, trường, người điều phối, số di động, email, nhân viên Sales', 'Cần bổ sung loại BEO và liên kết đoàn đã chốt'),
  ('Ngày và số lượng', 'Ngày đến, giờ khách đến, tổng khách, số hướng dẫn viên và điều hành, tài xế và phụ xe, đối tượng Corp hoặc TA', 'Cần bổ sung các số lượng riêng và giờ thực hiện'),
  ('Nhà hàng và thực đơn', 'Vị trí ăn, loại hình F&B, Bento, set menu mặn hoặc chay, giá theo khách hoặc bàn, số khách, số bàn, giờ ăn', 'Cần bổ sung thông tin ăn uống chi tiết'),
  ('Chuẩn bị phục vụ', 'Bàn giáo viên, suất tài xế, suất hướng dẫn viên, địa điểm và giờ giao suất ăn, bàn ghế, dụng cụ, nhân sự', 'Cần bổ sung nhiệm vụ và thông tin phục vụ'),
  ('Vận hành và cảnh quan', 'Số kệ, lộ trình tham quan, khu vực vệ sinh, yêu cầu vận hành công viên', 'Cần bổ sung nội dung giao bộ phận'),
  ('Kỹ thuật và an ninh', 'Yêu cầu điện nước, số xe và số chỗ, bố trí giữ xe, điều phối xe và an ninh', 'Cần bổ sung nội dung giao bộ phận'),
  ('Soát vé và IT', 'Giờ đến, hình thức soát vé, số lượng vé QR cần chuẩn bị, chương trình phối hợp', 'Cần bổ sung nội dung điều phối; mẫu Word không tự phát hành vé QR'),
  ('Marketing và kế toán', 'Giờ xác nhận trưởng đoàn, nội dung màn hình chào mừng, ấn phẩm, số vé, doanh thu, hóa đơn, tiền cọc và ngày cọc', 'Cần bổ sung nhiệm vụ và liên kết tài chính'),
  ('Duyệt và hoàn thành', 'Nhân viên Sales, Phòng KD, quản lý vận hành, bếp trưởng; thời gian xác nhận và trạng thái từng bộ phận', 'Cần bổ sung người duyệt và trạng thái thực hiện; tên chức danh có thể giữ cố định'),
 ]),
 ('Hợp đồng tham quan Sơn Tiên', [
  ('Thông tin hợp đồng và bên đặt', '{{ma_hop_dong}} / {{ngay_hop_dong}} / {{ten_khach_hang}} / {{ma_so_thue}} / {{dia_chi_cong_ty}}', 'Đã hỗ trợ trong Hợp đồng'),
  ('Đại diện bên đặt', '{{nguoi_dai_dien_phap_luat}} / {{chuc_vu_dai_dien}}', 'Đã hỗ trợ trong Hợp đồng'),
  ('Giá trị và thanh toán', '{{gia_tri_hop_dong}} / {{gia_tri_bang_chu}} / {{moc_thanh_toan}}', 'Đã hỗ trợ trong Hợp đồng; mốc là nội dung tổng hợp, chưa có mã riêng cho từng khoản cọc'),
  ('Lịch và chương trình', 'Ngày giờ tham quan, giờ ăn, nhà hàng, thực đơn, số bàn và số khách', 'Cần bổ sung liên kết thông tin đoàn vào xuất hợp đồng'),
  ('Vé và bảng giá dịch vụ', 'Đối tượng, chiều cao, đơn giá vé, suất miễn; nhóm vé, ăn uống, dịch vụ khác và tổng theo nhóm', 'Cần bổ sung bảng và quy tắc riêng; Hợp đồng hiện chỉ công bố mã dòng hàng cơ bản'),
  ('Thanh toán từng lần', 'Tỷ lệ, số tiền bằng số và chữ, ngày hạn, số thực thu, số còn lại và phát sinh', 'Cần bổ sung mã theo mốc và liên kết thu tiền'),
 ]),
 ('Hợp đồng liên kết du lịch', [
  ('Thông tin và hiệu lực', '{{ma_hop_dong}} / {{ngay_hop_dong}} / {{ngay_hieu_luc}} / {{ngay_het_han}}', 'Đã hỗ trợ trong Hợp đồng'),
  ('Đối tác và tham chiếu', '{{ten_khach_hang}} / {{email}} / {{ma_so_thue}} / {{ma_bao_gia}} / {{ma_don_hang}}', 'Đã hỗ trợ trong Hợp đồng'),
  ('Ưu đãi và miễn vé', 'Loại vé, đối tượng, giá hoặc tỷ lệ ưu đãi, tiêu chuẩn giáo viên, hướng dẫn viên và tài xế', 'Cần bổ sung bảng chính sách có hiệu lực theo thời kỳ'),
  ('Đặt đoàn theo hợp đồng', 'Danh sách phiếu đặt đoàn, ngày đi, số khách, tổng tiền, cọc và còn lại từng đoàn', 'Cần bổ sung quan hệ một hợp đồng nhiều đoàn'),
  ('Gia hạn và phụ lục', 'Chu kỳ gia hạn, phụ lục áp dụng, phiên bản và ngày hiệu lực chính sách', 'Cần bổ sung trường quản lý riêng; điều khoản có thể giữ cố định'),
 ]),
 ('Hợp đồng liên kết Resort', [
  ('Thông tin hai bên và hợp đồng', '{{ma_hop_dong}} / {{ngay_hop_dong}} / {{ten_khach_hang}} / {{so_dien_thoai}} / {{email}}', 'Đã hỗ trợ trong Hợp đồng; thông tin bên cung cấp đặt cố định hoặc cấu hình riêng'),
  ('Bảng giá phòng và combo', 'Hạng phòng, giá trong tuần, cuối tuần, lễ Tết; giá phòng và giá combo', 'Cần bổ sung bảng giá lưu trú nhiều cột và hiệu lực áp dụng'),
  ('Thông tin đặt phòng', 'Mã đặt, ngày nhận và trả phòng, số đêm, số phòng, số khách, người lớn, trẻ em, giường phụ và phụ thu', 'Cần bổ sung hồ sơ đặt phòng và mã trường tương ứng'),
  ('Xác nhận và thanh toán', 'Người đặt, yêu cầu đặc biệt, tiền cọc, hạn thanh toán, đổi hủy và phát sinh', 'Cần bổ sung liên kết đặt phòng và tài chính'),
  ('Quyền lợi và phụ lục', 'Quyền lợi phòng, combo, bữa sáng, dịch vụ kèm và phụ lục giá', 'Nội dung cố định trong Word hoặc danh mục chính sách nếu cần tự động'),
 ]),
 ('Hợp đồng nguyên tắc OTA', [
  ('Thông tin hợp đồng và đối tác', '{{ma_hop_dong}} / {{ngay_hop_dong}} / {{ten_khach_hang}} / {{dia_chi_cong_ty}} / {{ma_so_thue}}', 'Đã hỗ trợ trong Hợp đồng'),
  ('Đầu mối và hiệu lực', 'Đầu mối hai bên, chức vụ, số di động, email, thời hạn và chu kỳ gia hạn', 'Khách hàng có liên hệ cơ bản; cần bổ sung đầu mối nhiều bên và chu kỳ'),
  ('Vé và chính sách đại lý', 'Khu vui chơi, hình thức vé, loại vé, giá công bố, giá đại lý, hiệu lực bảng giá', 'Cần bổ sung bảng giá đối tác và phiên bản chính sách'),
  ('Yêu cầu và xác nhận dịch vụ', 'Số lượng vé, ngày sử dụng, ngày nhận, người nhận, hạn dùng vé và mã xác nhận', 'Cần bổ sung hồ sơ đặt vé; không tự phát hành vé từ trường trộn'),
  ('Phụ lục và thanh toán', 'Số phụ lục, ngày hiệu lực, hợp đồng gốc, ký quỹ, số tiền thanh toán và đối chiếu hoàn trả', 'Cần bổ sung liên kết phụ lục và dữ liệu đối soát'),
 ]),
 ('Nghiệm thu mua hàng vật tư và dịch vụ', [
  ('Nhận diện biên bản', 'Số biên bản, mã biểu mẫu, lần soát xét, ngày hiệu lực, ngày nghiệm thu, giờ bắt đầu và kết thúc', 'Cần bổ sung loại nghiệm thu mua hàng; mã biểu mẫu có thể giữ cố định'),
  ('Căn cứ mua hàng', 'Số và ngày phiếu đề nghị cấp mua, đơn vị yêu cầu, đơn mua và nhà cung cấp', 'Cần nối luồng mua hàng; không dùng khách hàng bán hàng thay nhà cung cấp'),
  ('Đại diện tham gia', 'Nhiều đại diện bên mua, bên bán, chức vụ và tổ nghiệm thu', 'Cần bổ sung danh sách đại diện'),
  ('Dòng kiểm tra', 'Tên hàng hoặc dịch vụ, đặc điểm quy cách, đơn vị tính, số theo chứng từ, số thực tế, tình trạng và ghi chú', 'Cần bổ sung số lượng thực tế, quy cách và kết quả theo dòng'),
  ('Kết quả và ký', 'Phương thức nghiệm thu, ý kiến, kết luận, đại diện hai bên, tài chính kế toán, cung ứng vật tư, thanh tra và đơn vị yêu cầu', 'Cần bổ sung kết quả và xác nhận; nghiệm thu giao khách hiện tại là loại khác'),
 ]),
 ('Báo cáo đề xuất công việc', [
  ('Nhận diện', 'Số đề xuất, mã biểu mẫu, lần ban hành, ngày ban hành, ngày đề xuất, khẩn hoặc thường', 'Cần bổ sung mẫu xuất đề xuất; mã biểu mẫu có thể giữ cố định'),
  ('Người gửi và người nhận', 'Người đề xuất, vị trí công tác, phòng ban, gửi đến, cấp có thẩm quyền', 'Cần nối hồ sơ người dùng và luồng duyệt'),
  ('Nội dung và duyệt', 'Nội dung đề xuất, người ký từng cấp, chức danh, ngày duyệt và ý kiến', 'Cần bổ sung trường của hồ sơ đề xuất và lịch sử duyệt'),
 ]),
]
for title, rows in sections:
    doc.add_page_break()
    doc.add_heading(title, 1)
    table(rows, widths=(4.1, 8.0, 6.1), headers=('Nhóm thông tin', 'Trường đang có hoặc dữ liệu cần điền', 'Khả năng sử dụng hiện tại'))
doc.add_page_break()
doc.add_heading('Lưu ý áp dụng mẫu Sơn Tiên', 1)
p('Phụ lục này là bảng đối chiếu nhu cầu biểu mẫu, không xác nhận rằng toàn bộ luồng Sơn Tiên đã được triển khai. Các mục Cần bổ sung phải có nơi nhập hoặc lấy dữ liệu, mã trường được công bố và kiểm tra xuất Word trước khi sử dụng tự động.')
p('Hai file Form mẫu BEO.doc và HĐNT - OTA 2026.doc cần chuyển thành .docx trước khi đưa vào cơ chế mẫu Word hiện tại. Không đổi phần mở rộng thay cho chuyển định dạng.')
p('Bộ mẫu có năm và thời điểm hiệu lực chưa thống nhất, đặc biệt ở phụ lục OTA. Cần xác nhận phiên bản áp dụng, thông tin bên cung cấp, tỷ lệ đặt cọc, điều kiện miễn vé và chính sách đổi hủy trước khi dùng chính thức.')
out = ROOT / 'resources/templates/print/bo-truong-tron.docx'
doc.save(out)
print(f'{out}\n{len(set.union(*(set(fs) for fs in sets.values())))} distinct fields; {len(sets)} document types')
